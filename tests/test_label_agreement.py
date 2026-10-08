"""v1.5 라벨 일치율 계산(agreement.py, triage-agreement) 테스트.

판정 수집 전에 계산 방법을 고정하기 위한 테스트다. 실제 판정 대신, 생성 도구로 만든 빈 시트에 답을 채워 넣어 확인한다.
"""

import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from triage_eval.pipeline import agreement as ag
from triage_eval.pipeline.gate import GateError

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools" / "label_agreement"))
import build_sheets as bs  # noqa: E402

COLS = {"분류": "I", "모듈1": "J", "모듈2 (선택)": "K", "우선순위": "L", "재현 정보": "M", "발생 빈도": "N", "처리": "O",
        "확신도": "P", "메모": "Q"}


@pytest.fixture
def crit():
    return ag.load_criteria()


@pytest.fixture
def blank(tmp_path):
    return bs.build(tmp_path / "blank", map_path=tmp_path / "map.json")


@pytest.fixture
def items(crit):
    return ag.load_items(crit)[0]


def label_answers(item: dict) -> dict:
    """정답 라벨의 최선 답으로 채운 판정."""
    lb = item["labels"]
    mods = lb["모듈"] or [bs.NO_MODULE]
    return {"분류": lb["분류"][0], "모듈1": mods[0], "모듈2 (선택)": mods[1] if len(mods) > 1 else None,
            "우선순위": lb["우선순위"][0], "재현 정보": lb["재현 정보"][0], "발생 빈도": lb["발생 빈도"][0],
            "처리": lb["처리"][0], "확신도": "높음", "메모": None}


def fill(src: Path, dst: Path, items: dict, change=None) -> Path:
    """빈 시트에 답을 채운다. change(번호, 답) 로 일부를 바꿀 수 있다."""
    wb = load_workbook(src)
    ws = wb["판정"]
    for r in range(2, ws.max_row + 1):
        q = ws[f"A{r}"].value
        ans = label_answers(items[q])
        if change:
            change(q, ans)
        for col, letter in COLS.items():
            ws[f"{letter}{r}"] = ans.get(col)
    wb.save(dst)
    return dst


def sheets(blank, tmp_path, items, change1=None, change2=None):
    return [fill(blank[0], tmp_path / "r1.xlsx", items, change1), fill(blank[1], tmp_path / "r2.xlsx", items, change2)]


# ── κ ───────────────────────────────────────────────────────

def test_cohen_kappa_textbook_value():
    """2×2 표 [[20, 5], [10, 15]] — po 0.7, pe 0.5 → κ 0.4."""
    a = ["예"] * 25 + ["아니오"] * 25
    b = ["예"] * 20 + ["아니오"] * 5 + ["예"] * 10 + ["아니오"] * 15
    assert ag.cohen_kappa(a, b) == pytest.approx(0.4)


def test_cohen_kappa_perfect_and_undefined():
    assert ag.cohen_kappa(["a", "b", "a"], ["a", "b", "a"]) == 1.0
    assert ag.cohen_kappa(["a", "a"], ["a", "a"]) is None   # 설계 의도 3 — 우연 일치 확률 1


# ── 전 과정 ─────────────────────────────────────────────────

def test_label_answers_are_fully_trusted(blank, tmp_path, crit, items):
    """정답 라벨 그대로 판정하면 모든 칸 일치율 100%, Critical 인식률 100%, 불일치 0건."""
    paths = sheets(blank, tmp_path, items)
    assert ag.run(paths, crit, write=False) == 0
    a = ag.read_sheet(paths[0], "채점자1")
    stats = ag.rater_stats(a, items)
    assert all(stats[f]["allowed_rate"] == 100.0 for f in ag.JUDGED)
    assert stats["critical"] == {"n": 20, "hit": 20, "recall": 100.0, "missed": []}
    assert ag.disagreements({"채점자1": a}, items, ag.load_items(crit)[1]) == []


def test_module2_is_joined_and_scored_like_model(blank, tmp_path, items):
    """설계 의도 1 — 모듈1·모듈2를 ' / '로 합쳐, 고른 모듈이 모두 허용 답 안에 있어야 일치."""
    multi = next(q for q, i in items.items() if len(i["labels"]["모듈"]) > 1)
    single = next(q for q, i in items.items() if len(i["labels"]["모듈"]) == 1)
    outside = next(m for m in bs.MODULES if m not in items[single]["labels"]["모듈"])

    def change(q, ans):
        if q == single:
            ans["모듈2 (선택)"] = outside   # 허용 답 밖 모듈을 하나 더 고름

    path = fill(blank[0], tmp_path / "m.xlsx", items, change)
    a = ag.read_sheet(path, "채점자1")
    assert " / " in a[multi]["모듈"]
    stats = ag.rater_stats(a, items)
    assert stats["모듈"]["allowed_ok"] == 62
    assert stats["모듈"]["kappa"] == 1.0   # κ는 모듈1 기준이라 영향 없음


def test_low_agreement_field_is_not_trusted(blank, tmp_path, crit, items):
    """채점자2가 개발 배정 문항을 모두 등록(재현 대기)으로 보내면 처리 칸이 기준에 못 미치고, 그 문항들이 검토표에 오른다."""
    def change2(q, ans):
        if ans["처리"] == "개발 배정":
            ans["처리"] = "등록(재현 대기)"

    paths = sheets(blank, tmp_path, items, change2=change2)
    assert ag.run(paths, crit, write=False) == 1
    a = {r: ag.read_sheet(p, r) for r, p in zip(crit["meta"]["raters"], paths)}
    stats = {r: ag.rater_stats(x, items) for r, x in a.items()}
    verdict = ag.judge(stats, crit)
    assert not verdict["fields"]["처리"]["trusted"] and verdict["fields"]["분류"]["trusted"]
    rows = ag.disagreements(a, items, ag.load_items(crit)[1])
    assert rows and all(x["칸"] == "처리" and x["허용 답 밖"] == ["채점자2"] and x["채점자끼리 다름"] for x in rows)


def test_missed_critical_fails_critical_recall(blank, tmp_path, crit, items):
    """Critical 문항 3건을 Major·등록(재현 대기)으로 보면 인식률 85% < 90%."""
    crit_qs = sorted(q for q, i in items.items() if "X-1" in i["labels"]["risk"])[:3]

    def change1(q, ans):
        if q in crit_qs:
            ans["우선순위"], ans["처리"] = "Major", "등록(재현 대기)"

    paths = sheets(blank, tmp_path, items, change1=change1)
    a = ag.read_sheet(paths[0], "채점자1")
    c = ag.rater_stats(a, items)["critical"]
    assert c["recall"] == 85.0 and c["missed"] == crit_qs
    assert ag.run(paths, crit, write=False) == 1


def test_report_has_no_model_output(blank, tmp_path, crit, items, monkeypatch):
    """설계 의도 7 — 검토표에 모델 응답이 들어가지 않는다 (실행 기록을 읽지 않는다)."""
    monkeypatch.setattr(ag, "RESULT_JSON", tmp_path / "out.json")
    monkeypatch.setattr(ag, "REPORT", tmp_path / "out.md")
    paths = sheets(blank, tmp_path, items, change2=lambda q, a: a.update(분류="버그 아님") if a["분류"] == "개선 제안" else None)
    ag.run(paths, crit, write=True)
    text = (tmp_path / "out.md").read_text(encoding="utf-8")
    assert "불일치 검토표" in text and "response" not in text and "[요약]" not in text
    saved = json.loads((tmp_path / "out.json").read_text(encoding="utf-8"))
    assert set(saved) == {"raters", "sheets", "stats", "between_raters_kappa", "verdict", "disagreements", "answers"}


# ── 계산 거부 (종료 코드 2) ─────────────────────────────────

def test_blank_cell_is_refused(blank, tmp_path, items):
    """설계 의도 4 — 빈칸은 조용히 오답·제외로 처리하지 않고 어디가 비었는지 알려 준다."""
    path = fill(blank[0], tmp_path / "b.xlsx", items, lambda q, a: a.update(처리=None) if q == "Q05" else None)
    with pytest.raises(GateError, match="Q05 처리 빈칸"):
        ag.read_sheet(path, "채점자1")


def test_value_outside_choices_is_refused(blank, tmp_path, items):
    path = fill(blank[0], tmp_path / "v.xlsx", items, lambda q, a: a.update(우선순위="중간") if q == "Q07" else None)
    with pytest.raises(GateError, match="선택지 밖 값 '중간'"):
        ag.read_sheet(path, "채점자1")


def test_swapped_sheets_are_refused(blank, tmp_path, crit, items):
    """설계 의도 5 — 채점자1 자리에 채점자2 시트를 넣으면 계산하지 않는다."""
    paths = sheets(blank, tmp_path, items)
    assert ag.main(["--sheet", str(paths[1]), "--sheet", str(paths[0]), "--no-write"]) == 2


def test_changed_dataset_is_refused(crit, monkeypatch, tmp_path):
    """설계 의도 6 — 대응표를 만든 뒤 평가셋이 바뀌면 계산하지 않는다."""
    changed = tmp_path / "ds.json"
    changed.write_text(ag.DATASET.read_text(encoding="utf-8") + " ", encoding="utf-8")
    monkeypatch.setattr(ag, "DATASET", changed)
    with pytest.raises(GateError, match="평가셋이 바뀌었습니다"):
        ag.load_items(crit)


def test_criteria_typo_is_rejected(tmp_path):
    bad = tmp_path / "c.toml"
    bad.write_text(ag.CRITERIA.read_text(encoding="utf-8").replace("recall_min", "recal_min"), encoding="utf-8")
    with pytest.raises(GateError):
        ag.load_criteria(bad)


def test_committed_criteria_match_plan(crit):
    """사전 등록한 기준 값 (계획 §4)."""
    f = crit["field"]
    assert all(f[k] == {"allowed_rate_min": 85.0, "kappa_min": 0.61} for k in ("처리", "분류", "우선순위"))
    assert all(f[k] == {"allowed_rate_min": 75.0} for k in ("모듈", "재현 정보", "발생 빈도"))
    assert crit["critical"]["recall_min"] == 90.0 and crit["meta"]["raters"] == ["채점자1", "채점자2"]

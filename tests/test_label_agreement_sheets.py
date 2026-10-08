"""v1.5 라벨 일치율 — 채점자용 판정 시트·안내서 테스트 (tools/label_agreement/).

채점자에게 정답이 새지 않는지, 모델이 받는 입력과 같은 내용이 들어가는지, 안내서가 기준서 규칙을 빠짐없이 옮겼는지 확인한다.
"""

import json
import re
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools" / "label_agreement"))
import build_sheets as bs  # noqa: E402

from triage_eval.pipeline.contract import ENUMS_V2  # noqa: E402

DATA = json.loads(bs.DATASET.read_text(encoding="utf-8"))
DEV = {i["id"]: i for i in DATA["items"] if i["split"] == "dev"}
GUIDE = (ROOT / "docs" / "label_agreement" / "rater_guide.md").read_text(encoding="utf-8")
GUIDELINE = (ROOT / "docs" / "dataset" / "triage_guideline.md").read_text(encoding="utf-8")
ITEM_ID = re.compile(r"(?<![A-Za-z0-9-])[ABSF]\d{2}[bc]?(?![0-9])")


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("sheets")
    map_path = out / "map.json"
    paths = bs.build(out, map_path=map_path)
    return paths, json.loads(map_path.read_text(encoding="utf-8"))


def test_blind_map_is_reproducible_and_matches_committed(built):
    """같은 seed면 같은 순서 — 저장소의 대응표와 다시 만든 대응표가 같다 (build_sheets.py 설계 의도 2)."""
    _, made = built
    committed = json.loads(bs.MAP_PATH.read_text(encoding="utf-8"))
    assert made == committed


def test_blind_map_covers_dev_only(built):
    _, made = built
    ids = list(made["map"].values())
    assert sorted(ids) == sorted(DEV) and len(ids) == 63
    assert list(made["map"]) == [f"Q{n:02d}" for n in range(1, 64)]
    assert ids != sorted(ids)   # 원래 순서(트랙·세트별)가 드러나지 않게 섞였다


def test_two_sheets_have_same_items(built):
    paths, _ = built
    tables = [[[c.value for c in row[:8]] for row in load_workbook(p)["판정"].iter_rows(min_row=2)] for p in paths]
    assert len(paths) == 2 and tables[0] == tables[1] and len(tables[0]) == 63


def test_sheet_has_model_input_only_and_empty_answers(built):
    """설계 의도 1 — 입력은 모델 입력과 같은 내용, 답 칸은 비어 있고, 정답·근거·원래 번호는 없다."""
    paths, made = built
    ws = load_workbook(paths[0])["판정"]
    header = [c.value for c in ws[1]]
    assert header == bs.INPUT_COLS + bs.ANSWER_COLS + bs.EXTRA_COLS
    for row in ws.iter_rows(min_row=2):
        values = dict(zip(header, (c.value for c in row)))
        item = DEV[made["map"][values["번호"]]]
        expected = bs.row_of(values["번호"], item)
        assert all((values[k] or "") == expected[k] for k in bs.INPUT_COLS)
        assert all(values[k] is None for k in bs.ANSWER_COLS + bs.EXTRA_COLS)
        text = " ".join(str(v) for v in values.values() if v)
        assert item["rationale"] not in text and item["id"] not in text


def test_dropdowns_use_contract_values(built):
    """설계 의도 4 — 선택지가 출력 형식 v2의 허용 값과 같아, 받은 답을 모델과 같은 채점 함수로 대조할 수 있다."""
    paths, _ = built
    wb = load_workbook(paths[0])
    choices = wb["선택지"]
    cols = {choices.cell(row=1, column=c).value: [choices.cell(row=r, column=c).value
                                                   for r in range(2, choices.max_row + 1)
                                                   if choices.cell(row=r, column=c).value]
            for c in range(1, len(bs.CHOICES) + 1)}
    for field in ("분류", "우선순위", "재현 정보", "발생 빈도", "처리"):
        assert cols[field] == ENUMS_V2[field]
    validated = {str(dv.sqref).split(":")[0][:1] for dv in wb["판정"].data_validations.dataValidation}
    assert validated == set("IJKLMNOP")   # 분류 ~ 확신도 8칸


def test_guide_has_no_eval_item_ids():
    """안내서에 평가 문항 번호가 없다 — 기준서의 경계 규칙 표에는 문항 번호가 적혀 있어 그대로 보내면 정답이 드러난다."""
    assert not ITEM_ID.findall(GUIDE)
    assert ITEM_ID.findall(GUIDELINE)   # 원본 기준서에는 있다 (확인 방법이 동작하는지)


def test_guide_carries_every_guideline_rule():
    """안내서가 기준서의 판정 규칙 번호를 빠짐없이 옮겼다."""
    codes = set(re.findall(r"\b(?:C-\d+b?|R-\d|F-\d|H-\d+[a-d]?|P-B\d)\b", GUIDELINE.split("## 7.")[0]))
    codes -= {"F-1", "F-2", "F-3"}   # 발생 빈도 규칙은 번호 대신 표와 문장으로 옮겼다 (아래 확인)
    missing = sorted(c for c in codes if c not in GUIDE)
    assert not missing, missing
    for phrase in ("드문 조건에서만 생기고 우회 경로가 있는 무료 보상 손실", "적힌", "조건문"):
        assert phrase in GUIDE
